"""Shared XLSX/CSV reader for the soft-coded master pages."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from openpyxl import load_workbook

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


def _row_limit_message(max_rows: int) -> str:
    return f"File exceeds the {max_rows:,}-row import limit."


def _byte_limit_message(max_bytes: int) -> str:
    return f"File exceeds the {max_bytes / (1024 * 1024):g} MB import limit."


def _check_stream_row_limit(row_number: int, max_rows: int | None) -> None:
    """Guard the streaming reader.

    ``max_rows`` counts data rows, so the header row is excluded. ``None`` (or
    ``0``) means the caller imposes no ceiling — calibration charts, which run
    to tens of thousands of points, are read that way.
    """
    if max_rows and row_number > max_rows + 1:
        raise ValueError(_row_limit_message(max_rows))


def normalise_header(value) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[\s\-]+", "_", text)
    return re.sub(r"[^a-z0-9_]", "", text)


def _read_table(
    content: bytes,
    suffix: str,
    *,
    max_rows: int | None = MAX_IMPORT_ROWS,
) -> list[list[str]]:
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                _check_stream_row_limit(row_number, max_rows)
                rows.append(["" if value is None else str(value).strip() for value in row])
            return rows
        finally:
            workbook.close()

    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        decoded = content.decode("latin-1")
    rows = []
    for row_number, row in enumerate(csv.reader(io.StringIO(decoded)), start=1):
        _check_stream_row_limit(row_number, max_rows)
        rows.append([str(value or "").strip() for value in row])
    return rows


def read_master_import(
    content: bytes,
    filename: str,
    *,
    max_rows: int | None = MAX_IMPORT_ROWS,
    max_bytes: int | None = MAX_IMPORT_BYTES,
) -> dict:
    """Read an upload into normalized row dictionaries without changing data.

    ``max_rows`` and ``max_bytes`` default to the master-data ceilings. Callers
    that legitimately handle far larger files — the calibration chart parser,
    whose books run to tens of thousands of points — pass their own ceiling, or
    ``None`` to read the file without one.
    """
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("Unsupported file type. Use .xlsx or .csv.")
    if max_bytes and len(content) > max_bytes:
        raise ValueError(_byte_limit_message(max_bytes))

    try:
        table = _read_table(content, suffix, max_rows=max_rows)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("The file could not be read. Check that it is a valid XLSX or CSV file.") from exc
    table = [row for row in table if any(str(cell or "").strip() for cell in row)]
    if not table:
        return {
            "filename": filename,
            "format": suffix.lstrip("."),
            "columns": [],
            "rows": [],
        }

    headers = [normalise_header(value) for value in table[0]]
    rows = []
    for row_number, values in enumerate(table[1:], start=2):
        rows.append(
            {
                "row_number": row_number,
                **{
                    header: str(values[index] or "").strip()
                    for index, header in enumerate(headers)
                    if header and index < len(values)
                },
            }
        )
    if max_rows and len(rows) > max_rows:
        raise ValueError(_row_limit_message(max_rows))
    return {
        "filename": filename,
        "format": suffix.lstrip("."),
        "columns": [header for header in headers if header],
        "rows": rows,
    }
