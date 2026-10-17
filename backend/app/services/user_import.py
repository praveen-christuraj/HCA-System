"""User import from XLSX/CSV with preview and row-level validation (rule 7).

Parsing is server-side (openpyxl for Excel, csv for CSV) so the browser bundle
stays free of spreadsheet libraries. Passwords are never echoed back in the
preview or written to the audit trail.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Role, User, UserRole
from app.utils.password_policy import password_policy_errors
from app.utils.security import hash_password
from app.utils.time import utcnow

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000

REQUIRED_COLUMNS = {"username", "full_name", "email", "password"}

HEADER_ALIASES = {
    "user": "username",
    "user_name": "username",
    "login": "username",
    "name": "full_name",
    "fullname": "full_name",
    "full_name": "full_name",
    "email_address": "email",
    "mail": "email",
    "mobile": "phone",
    "phone_number": "phone",
    "dept": "department",
    "role": "role_names",
    "roles": "role_names",
    "role_name": "role_names",
    "force_change": "force_password_change",
    "force_password_change": "force_password_change",
}

USERNAME_PATTERN = re.compile(r"^[a-zA-Z0-9._-]+$")
SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


def _normalise_header(value) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[\s\-]+", "_", text)
    text = re.sub(r"[^a-z0-9_]", "", text)
    return HEADER_ALIASES.get(text, text)


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _read_table(content: bytes, suffix: str) -> list[list[str]]:
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [
            [_cell_text(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        workbook.close()
        return rows

    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    return [[_cell_text(cell) for cell in row] for row in csv.reader(io.StringIO(text))]


def _parse_bool(value: str, default: bool = True) -> bool:
    if not value:
        return default
    return value.strip().lower() in {"1", "true", "yes", "y", "on"}


def parse_user_import(content: bytes, filename: str, db: Session) -> dict:
    suffix = Path(filename or "").suffix.lower()
    payload = {
        "filename": filename,
        "format": suffix.lstrip("."),
        "total_rows": 0,
        "valid_rows": 0,
        "invalid_rows": 0,
        "columns_detected": [],
        "missing_columns": [],
        "rows": [],
        "message": "",
    }

    if suffix not in SUPPORTED_SUFFIXES:
        payload["message"] = "Unsupported file type. Use .xlsx or .csv."
        return payload
    if len(content) > MAX_IMPORT_BYTES:
        payload["message"] = "File exceeds the 5 MB import limit."
        return payload

    try:
        table = _read_table(content, suffix)
    except Exception:
        payload["message"] = "The file could not be read. Check that it is a valid XLSX or CSV file."
        return payload

    table = [row for row in table if any(cell for cell in row)]
    if not table:
        payload["message"] = "The file is empty."
        return payload

    raw_headers = table[0]
    headers = [_normalise_header(cell) for cell in raw_headers]
    payload["columns_detected"] = [header for header in headers if header]
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    payload["missing_columns"] = missing
    if missing:
        payload["message"] = "Missing required columns: " + ", ".join(missing)
        return payload

    index_by_header = {}
    for index, header in enumerate(headers):
        if header and header not in index_by_header:
            index_by_header[header] = index

    existing_usernames = {
        str(value).strip().lower()
        for (value,) in db.query(User.username).all()
    }
    existing_emails = {
        str(value).strip().lower()
        for (value,) in db.query(User.email).all()
    }
    roles = db.query(Role).filter(Role.status == "Active").all()
    roles_by_name = {role.role_name.strip().lower(): role for role in roles}

    seen_usernames: set[str] = set()
    seen_emails: set[str] = set()
    preview_rows = []

    data_rows = table[1:]
    if len(data_rows) > MAX_IMPORT_ROWS:
        payload["message"] = f"File exceeds the {MAX_IMPORT_ROWS}-row import limit."
        return payload

    for offset, row in enumerate(data_rows, start=2):
        def value(name: str) -> str:
            index = index_by_header.get(name)
            if index is None or index >= len(row):
                return ""
            return row[index].strip()

        username = value("username")
        full_name = value("full_name")
        email = value("email")
        password = value("password")
        status = value("status") or "Active"

        role_names_raw = value("role_names")
        role_names = [
            name.strip()
            for name in re.split(r"[;,]", role_names_raw)
            if name.strip()
        ]

        errors: list[str] = []
        if not username:
            errors.append("Username is required")
        elif len(username) < 3 or len(username) > 80:
            errors.append("Username must be between 3 and 80 characters")
        elif not USERNAME_PATTERN.match(username):
            errors.append("Username may only contain letters, digits, dot, underscore or hyphen")

        normalised_username = username.lower()
        if normalised_username and normalised_username in existing_usernames:
            errors.append("Username already exists")
        if normalised_username and normalised_username in seen_usernames:
            errors.append("Duplicate username in file")
        if normalised_username:
            seen_usernames.add(normalised_username)

        if not full_name:
            errors.append("Full name is required")

        normalised_email = email.lower()
        if not email:
            errors.append("Email is required")
        else:
            try:
                normalised_email = validate_email(
                    email, check_deliverability=False
                ).normalized.lower()
            except EmailNotValidError as exc:
                errors.append(f"Invalid email: {exc}")
        if normalised_email and normalised_email in existing_emails:
            errors.append("Email already exists")
        if normalised_email and normalised_email in seen_emails:
            errors.append("Duplicate email in file")
        if normalised_email:
            seen_emails.add(normalised_email)

        if not password:
            errors.append("Password is required")
        else:
            errors.extend(password_policy_errors(password))

        if status not in {"Active", "Inactive"}:
            errors.append("Status must be Active or Inactive")

        resolved_roles = []
        for role_name in role_names:
            role = roles_by_name.get(role_name.lower())
            if role is None:
                errors.append(f"Unknown or inactive role: {role_name}")
            else:
                resolved_roles.append(role)

        preview_rows.append(
            {
                "row_number": offset,
                "username": username,
                "full_name": full_name,
                "email": normalised_email or email,
                "department": value("department"),
                "designation": value("designation"),
                "phone": value("phone"),
                "status": status,
                "role_names": ", ".join(role.role_name for role in resolved_roles),
                "valid": len(errors) == 0,
                "errors": errors,
                # Internal fields used only when the import is committed.
                "_password": password,
                "_role_ids": [role.id for role in resolved_roles],
                "_force_change": _parse_bool(value("force_password_change"), True),
                "_email": normalised_email or email,
            }
        )

    valid_rows = sum(1 for row in preview_rows if row["valid"])
    payload.update(
        {
            "total_rows": len(preview_rows),
            "valid_rows": valid_rows,
            "invalid_rows": len(preview_rows) - valid_rows,
            "rows": [
                {key: value for key, value in row.items() if not key.startswith("_")}
                for row in preview_rows
            ],
            "message": (
                f"{valid_rows} row(s) ready to import"
                if valid_rows
                else "No valid rows to import"
            ),
        }
    )
    payload["_parsed_rows"] = preview_rows
    return payload


def commit_user_import(db: Session, parsed: dict) -> int:
    created = 0
    for row in parsed.get("_parsed_rows", []):
        if not row["valid"]:
            continue
        user = User(
            full_name=row["full_name"],
            username=row["username"],
            email=row["_email"],
            phone=row["phone"] or None,
            department=row["department"] or None,
            designation=row["designation"] or None,
            password_hash=hash_password(row["_password"]),
            password_changed_at=utcnow(),
            force_password_change=row["_force_change"],
            status=row["status"],
            failed_login_count=0,
            totp_enabled=False,
            force_2fa=False,
            created_at=utcnow(),
            updated_at=utcnow(),
        )
        db.add(user)
        db.flush()
        for role_id in row["_role_ids"]:
            db.add(UserRole(user_id=user.id, role_id=role_id))
        created += 1
    return created
