"""Shared XLSX / CSV / PDF export builders (owner rule 7).

All exports use the same header treatment, timestamp and file naming. CSV
cells are sanitised against spreadsheet formula injection.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime

from fastapi import Response
from fastapi.responses import StreamingResponse

from app.utils.helpers import format_datetime
from app.utils.time import utcnow


def safe_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return format_datetime(value)
    text = str(value)
    if text[:1] in {"=", "+", "-", "@", "\t", "\r"}:
        return "'" + text
    return text


def timestamped_filename(prefix: str, extension: str) -> str:
    return f"{prefix}-{utcnow().strftime('%Y%m%d-%H%M%S')}.{extension}"


def csv_response(prefix: str, headers: list[str], rows: list[list]) -> Response:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([safe_cell(value) for value in row])
    filename = timestamped_filename(prefix, "csv")
    return Response(
        content="\ufeff" + buffer.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def xlsx_response(
    prefix: str,
    title: str,
    headers: list[str],
    rows: list[list],
    widths: list[int] | None = None,
) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31] or "Export"

    header_fill = PatternFill("solid", fgColor="12315F")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    if headers:
        last_column = sheet.cell(row=1, column=len(headers)).column_letter
        sheet.auto_filter.ref = f"A1:{last_column}{max(len(rows) + 1, 1)}"

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=safe_cell(value))
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical="top")

    if widths:
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = timestamped_filename(prefix, "xlsx")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _pdf_safe(value: str) -> str:
    return str(value).encode("latin-1", "replace").decode("latin-1")


def pdf_response(
    prefix: str,
    title: str,
    headers: list[str],
    rows: list[list],
    widths: list[int] | None = None,
    filters_summary: str = "no filters",
    orientation: str = "L",
    subtitle: str = "",
) -> Response:
    from fpdf import FPDF

    if not widths:
        col_w = max(16, int(270 / max(len(headers), 1)))
        widths = [col_w] * len(headers)

    pdf = FPDF(orientation=orientation, unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 14)
    pdf.set_text_color(18, 49, 95)
    pdf.cell(0, 8, _pdf_safe(f"HCA System - {title}"), new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(70, 85, 105)
    header_line = f"Generated: {format_datetime(utcnow())}  |  Rows: {len(rows)}"
    if subtitle:
        header_line += f"  |  {subtitle}"
    pdf.cell(0, 5, _pdf_safe(header_line), new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 4.5, _pdf_safe(f"Filters: {filters_summary}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    pdf.set_font("Helvetica", "B", 7.5)
    pdf.set_fill_color(18, 49, 95)
    pdf.set_text_color(255, 255, 255)
    for label, width in zip(headers, widths):
        pdf.cell(width, 6, _pdf_safe(label), border=0, fill=True)
    pdf.ln(6)

    pdf.set_font("Helvetica", "", 7)
    pdf.set_text_color(25, 35, 50)
    for row in rows:
        for width, value in zip(widths, row):
            text = _pdf_safe(safe_cell(value))
            max_chars = max(int(width / 1.55), 6)
            if len(text) > max_chars:
                text = text[: max_chars - 1] + "..."
            pdf.cell(width, 5, text, border="B")
        pdf.ln(5)

    filename = timestamped_filename(prefix, "pdf")
    return Response(
        content=bytes(pdf.output()),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
