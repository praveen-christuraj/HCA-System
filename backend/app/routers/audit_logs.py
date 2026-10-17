"""Audit Trail API.

Read-only audit list, exports (CSV/XLSX/PDF) and the persisted UI action
endpoint. Audit details are never stored as JSON (owner rule 4); the UI event
payload is flattened into a human-readable summary.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time

from fastapi import APIRouter, Depends, Query, Request, Response
from fastapi.responses import StreamingResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.config import MAX_EXPORT_ROWS
from app.database import get_db
from app.dependencies.auth import get_current_user
from app.dependencies.permissions import require_permission
from app.models import AuditLog, User
from app.schemas import AuditUiEventRequest
from app.services.audit_service import create_audit_log
from app.services.serializers import audit_to_dict
from app.utils.helpers import format_datetime
from app.utils.pagination import paginate_query
from app.utils.time import utcnow

router = APIRouter(tags=["Auditing"])

EXPORT_COLUMNS = [
    ("created_at", "Timestamp"),
    ("actor_username", "Actor"),
    ("actor_full_name", "Actor Name"),
    ("module_name", "Module"),
    ("action", "Action"),
    ("category", "Category"),
    ("outcome", "Outcome"),
    ("entity_type", "Entity Type"),
    ("entity_id", "Entity ID"),
    ("entity_label", "Entity"),
    ("http_method", "Method"),
    ("request_path", "Path"),
    ("status_code", "Status"),
    ("ip_address", "IP Address"),
    ("old_status", "Old Status"),
    ("new_status", "New Status"),
    ("remarks", "Remarks"),
    ("details_summary", "Details"),
]


def _filtered_query(
    db: Session,
    *,
    search: str | None,
    module_name: str | None,
    action: str | None,
    category: str | None,
    outcome: str | None,
    actor_username: str | None,
    date_from: date | None,
    date_to: date | None,
):
    query = db.query(AuditLog)
    if search:
        term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                AuditLog.actor_username.ilike(term),
                AuditLog.actor_full_name.ilike(term),
                AuditLog.module_name.ilike(term),
                AuditLog.action.ilike(term),
                AuditLog.entity_label.ilike(term),
                AuditLog.remarks.ilike(term),
                AuditLog.details_summary.ilike(term),
                AuditLog.request_path.ilike(term),
            )
        )
    if module_name:
        query = query.filter(AuditLog.module_name == module_name)
    if action:
        query = query.filter(AuditLog.action == action)
    if category:
        query = query.filter(AuditLog.category == category)
    if outcome:
        query = query.filter(AuditLog.outcome == outcome)
    if actor_username:
        query = query.filter(AuditLog.actor_username.ilike(f"%{actor_username.strip()}%"))
    # Dedicated date/date-range group (owner rule 6).
    if date_from:
        query = query.filter(AuditLog.created_at >= datetime.combine(date_from, time.min))
    if date_to:
        query = query.filter(AuditLog.created_at <= datetime.combine(date_to, time.max))
    return query


def _filters_summary(**filters) -> str:
    parts = [f"{key}={value}" for key, value in filters.items() if value]
    return "; ".join(parts) if parts else "no filters"


@router.get("/audit-logs")
def list_audit_logs(
    page: int = Query(1, ge=1),
    page_size: int = Query(20),
    search: str | None = None,
    module_name: str | None = None,
    action: str | None = None,
    category: str | None = None,
    outcome: str | None = None,
    actor_username: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    user: User = Depends(require_permission("audit:read")),
    db: Session = Depends(get_db),
):
    query = _filtered_query(
        db,
        search=search,
        module_name=module_name,
        action=action,
        category=category,
        outcome=outcome,
        actor_username=actor_username,
        date_from=date_from,
        date_to=date_to,
    )
    result = paginate_query(query.order_by(AuditLog.created_at.desc(), AuditLog.id.desc()), page, page_size)
    result["items"] = [audit_to_dict(entry) for entry in result["items"]]
    result["modules"] = [
        value
        for (value,) in db.query(AuditLog.module_name).distinct().order_by(AuditLog.module_name).all()
    ]
    result["categories"] = [
        value
        for (value,) in db.query(AuditLog.category).distinct().order_by(AuditLog.category).all()
    ]
    result["actions"] = [
        value
        for (value,) in db.query(AuditLog.action).distinct().order_by(AuditLog.action).all()
    ]
    return result


@router.post("/audit/ui-events", status_code=201)
def record_ui_event(
    payload: AuditUiEventRequest,
    request: Request,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Persist client-side actions (page visits, filters, theme, print)."""
    summary_parts = [f"route={payload.route or request.headers.get('referer', '')}"]
    for key, value in list((payload.details or {}).items())[:15]:
        text = str(value)
        summary_parts.append(f"{key}={text[:120]}")
    summary = "; ".join(summary_parts)[:1000]

    category = "Navigation" if payload.action.strip().lower() in {"open page", "navigate"} else "UI Activity"

    create_audit_log(
        db,
        module_name=payload.module,
        action=payload.action,
        category=category,
        current_user=user,
        request=request,
        remarks="Client-side action recorded for completeness",
        details_summary=summary,
        status_code=201,
    )
    db.commit()
    return {"message": "UI event recorded"}


def _safe_cell(value) -> str:
    """Block spreadsheet formula injection (owner rule 7)."""
    if value is None:
        return ""
    text = str(value)
    if text[:1] in {"=", "+", "-", "@", "\t", "\r"}:
        return "'" + text
    return text


def _export_rows(db: Session, filters: dict) -> list[AuditLog]:
    query = _filtered_query(db, **filters).order_by(AuditLog.created_at.asc(), AuditLog.id.asc())
    return query.limit(MAX_EXPORT_ROWS).all()


def _row_values(entry: AuditLog) -> list[str]:
    values = []
    for field, _label in EXPORT_COLUMNS:
        value = getattr(entry, field, None)
        if isinstance(value, datetime):
            value = format_datetime(value)
        values.append("" if value is None else str(value))
    return values


def _csv_response(rows: list[AuditLog]) -> Response:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([label for _field, label in EXPORT_COLUMNS])
    for entry in rows:
        writer.writerow([_safe_cell(value) for value in _row_values(entry)])

    filename = f"audit-log-{utcnow().strftime('%Y%m%d-%H%M%S')}.csv"
    return Response(
        content="\ufeff" + buffer.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _xlsx_response(rows: list[AuditLog]) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Audit Log"

    header_fill = PatternFill("solid", fgColor="12315F")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    for index, (_field, label) in enumerate(EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:R{max(len(rows) + 1, 1)}"

    for row_index, entry in enumerate(rows, start=2):
        for column_index, value in enumerate(_row_values(entry), start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_safe_cell(value))
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=column_index in {17, 18})

    widths = [20, 18, 24, 20, 26, 16, 12, 14, 10, 26, 10, 28, 8, 15, 14, 14, 36, 50]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    filename = f"audit-log-{utcnow().strftime('%Y%m%d-%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _pdf_safe(value: str) -> str:
    return str(value).encode("latin-1", "replace").decode("latin-1")


def _pdf_response(rows: list[AuditLog], filters: dict) -> Response:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 14)
    pdf.set_text_color(18, 49, 95)
    pdf.cell(0, 8, "HCA System - Audit Trail", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(70, 85, 105)
    pdf.cell(0, 5, _pdf_safe(f"Generated: {format_datetime(utcnow())}  |  Rows: {len(rows)}"), new_x="LMARGIN", new_y="NEXT")
    active_filters = _filters_summary(**filters)
    pdf.multi_cell(0, 4.5, _pdf_safe(f"Filters: {active_filters}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    columns = [
        ("Timestamp", 33),
        ("Actor", 33),
        ("Module", 36),
        ("Action", 38),
        ("Category", 28),
        ("Outcome", 16),
        ("Entity", 36),
        ("Details", 57),
    ]

    pdf.set_font("Helvetica", "B", 7.5)
    pdf.set_fill_color(18, 49, 95)
    pdf.set_text_color(255, 255, 255)
    for label, width in columns:
        pdf.cell(width, 6, label, border=0, fill=True)
    pdf.ln(6)

    pdf.set_text_color(25, 35, 50)
    pdf.set_font("Helvetica", "", 7)
    for entry in rows:
        values = [
            format_datetime(entry.created_at),
            entry.actor_username or "system",
            entry.module_name,
            entry.action,
            entry.category,
            entry.outcome,
            entry.entity_label or (entry.entity_type or ""),
            entry.details_summary or entry.remarks or "",
        ]
        for (label, width), value in zip(columns, values):
            text = _pdf_safe(value or "")
            max_chars = max(int(width / 1.55), 6)
            if len(text) > max_chars:
                text = text[: max_chars - 1] + "..."
            pdf.cell(width, 5, text, border="B")
        pdf.ln(5)

    content = bytes(pdf.output())
    filename = f"audit-log-{utcnow().strftime('%Y%m%d-%H%M%S')}.pdf"
    return Response(
        content=content,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/audit-logs/export")
def export_audit_logs(
    request: Request,
    format: str = Query("csv", pattern="^(csv|xlsx|pdf)$"),
    search: str | None = None,
    module_name: str | None = None,
    action: str | None = None,
    category: str | None = None,
    outcome: str | None = None,
    actor_username: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    user: User = Depends(require_permission("audit:read")),
    db: Session = Depends(get_db),
):
    filters = {
        "search": search,
        "module_name": module_name,
        "action": action,
        "category": category,
        "outcome": outcome,
        "actor_username": actor_username,
        "date_from": date_from,
        "date_to": date_to,
    }
    rows = _export_rows(db, filters)

    create_audit_log(
        db,
        module_name="Audit Trail",
        action="Export Audit Log",
        category="Export",
        current_user=user,
        request=request,
        remarks=f"Audit trail exported as {format.upper()}",
        details_summary=f"format={format}; rows={len(rows)}; {_filters_summary(**filters)}",
    )
    db.commit()

    if format == "xlsx":
        return _xlsx_response(rows)
    if format == "pdf":
        return _pdf_response(rows, filters)
    return _csv_response(rows)
